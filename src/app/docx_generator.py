from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import io
import re

def aplicar_fondo_celda(celda, color_hex):
    """Aplica color de fondo a una celda usando XML"""
    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    )
    tcPr = celda._element.tcPr
    if tcPr is None:
        tcPr = parse_xml(f'<w:tcPr {nsdecls("w")}/>')
        celda._element.insert(0, tcPr)
    tcPr.append(shading)

def crear_documento_unidad_didactica(contenido_raw, contenido_normalizado, area_curricular, titulo_unidad, sesiones_data, competencias_trans_data, propositos_data, valores_enfoques_data, situacion_significativa_texto):
    """
    Genera un documento DOCX programáticamente para la Unidad Didáctica,
    asegurando que las tablas complejas se formen correctamente.
    """
    doc = Document()
    
    # Configurar propiedades del documento
    doc.core_properties.title = titulo_unidad
    doc.core_properties.author = "Sistema IA Educativa"
    doc.core_properties.subject = titulo_unidad
    
    # Título principal
    titulo_principal = doc.add_heading("UNIDAD DIDÁCTICA", 0)
    titulo_principal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Área curricular
    area_heading = doc.add_paragraph()
    area_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    area_run = area_heading.add_run(f"ÁREA CURRICULAR: {area_curricular}")
    area_run.bold = True
    area_run.font.size = Pt(14)
    
    doc.add_paragraph() # break
    
    color_fondo_hex = "D9E1F2"
    
    # I. DATOS INFORMATIVOS
    doc.add_heading("I. DATOS INFORMATIVOS", 2)
    tabla_datos = doc.add_table(rows=7, cols=4)
    tabla_datos.style = 'Table Grid'
    
    for row in tabla_datos.rows:
        for cell in row.cells:
            celda_p = cell.paragraphs[0]
            celda_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Fila 1
    celda_dre = tabla_datos.cell(0, 0)
    celda_dre.text = "DRE / UGEL"
    celda_dre.paragraphs[0].runs[0].bold = True
    aplicar_fondo_celda(celda_dre, color_fondo_hex)
    
    celda_inst = tabla_datos.cell(0, 2)
    celda_inst.text = "Institución educativa"
    celda_inst.paragraphs[0].runs[0].bold = True
    aplicar_fondo_celda(celda_inst, color_fondo_hex)
    
    # Fila 2
    celda_area = tabla_datos.cell(1, 0)
    celda_area.text = "Área curricular"
    celda_area.paragraphs[0].runs[0].bold = True
    aplicar_fondo_celda(celda_area, color_fondo_hex)
    
    celda_grado = tabla_datos.cell(1, 2)
    celda_grado.text = "Grado y secciones"
    celda_grado.paragraphs[0].runs[0].bold = True
    aplicar_fondo_celda(celda_grado, color_fondo_hex)
    
    # Fila 3
    celda_fini = tabla_datos.cell(2, 0)
    celda_fini.text = "Fecha de inicio"
    celda_fini.paragraphs[0].runs[0].bold = True
    aplicar_fondo_celda(celda_fini, color_fondo_hex)
    
    celda_fterm = tabla_datos.cell(2, 2)
    celda_fterm.text = "Fecha de término"
    celda_fterm.paragraphs[0].runs[0].bold = True
    aplicar_fondo_celda(celda_fterm, color_fondo_hex)
    
    # Fila 4 (fusionada)
    celda_coord = tabla_datos.cell(3, 0)
    celda_coord.merge(tabla_datos.cell(4, 0))
    celda_coord.text = "Coordinación pedagógica"
    celda_coord.paragraphs[0].runs[0].bold = True
    aplicar_fondo_celda(celda_coord, color_fondo_hex)
    
    celda_coord_val1 = tabla_datos.cell(3, 1)
    celda_coord_val1.merge(tabla_datos.cell(3, 3))
    
    celda_coord_val2 = tabla_datos.cell(4, 1)
    celda_coord_val2.merge(tabla_datos.cell(4, 3))
    
    # Fila 6 (Docentes)
    celda_docentes = tabla_datos.cell(5, 0)
    celda_docentes.text = "Docente(s)"
    celda_docentes.paragraphs[0].runs[0].bold = True
    aplicar_fondo_celda(celda_docentes, color_fondo_hex)
    
    celda_doc_val = tabla_datos.cell(5, 1)
    celda_doc_val.merge(tabla_datos.cell(5, 3))

    doc.add_paragraph()
    
    # II. SITUACIÓN SIGNIFICATIVA
    doc.add_heading("II. SITUACIÓN SIGNIFICATIVA", 2)
    doc.add_paragraph(situacion_significativa_texto if situacion_significativa_texto else "[El contenido de la situación significativa se generará aquí]")
    
    doc.add_paragraph()
    
    # III. PROPÓSITOS DE APRENDIZAJE
    doc.add_heading("III. PROPÓSITOS DE APRENDIZAJE", 2)
    
    # Tabla Competencias transversales
    tabla_transv = doc.add_table(rows=3, cols=3)
    tabla_transv.style = 'Table Grid'
    hdr_cells = tabla_transv.rows[0].cells
    hdr_cells[0].text = "Competencias transversales"
    hdr_cells[1].text = "Estándares de aprendizaje"
    hdr_cells[2].text = "Instrumento"
    for c in hdr_cells:
        c.paragraphs[0].runs[0].bold = True
        aplicar_fondo_celda(c, color_fondo_hex)
        
    tabla_transv.cell(1, 0).text = "Se desenvuelve en los entornos virtuales generados por las TIC."
    tabla_transv.cell(1, 1).text = competencias_trans_data.get("Se desenvuelve en los entornos virtuales generados por las TIC.", {}).get("estandar", "")
    tabla_transv.cell(1, 2).text = competencias_trans_data.get("Se desenvuelve en los entornos virtuales generados por las TIC.", {}).get("instrumento", "")
    
    tabla_transv.cell(2, 0).text = "Gestiona su aprendizaje de manera autónoma."
    tabla_transv.cell(2, 1).text = competencias_trans_data.get("Gestiona su aprendizaje de manera autónoma.", {}).get("estandar", "")
    tabla_transv.cell(2, 2).text = competencias_trans_data.get("Gestiona su aprendizaje de manera autónoma.", {}).get("instrumento", "")
    
    doc.add_paragraph()

    # Tabla Competencias de área
    tabla_area = doc.add_table(rows=2, cols=6)
    tabla_area.style = 'Table Grid'
    hdr_cells = tabla_area.rows[0].cells
    hdr_cells[0].text = "Competencias de área"
    hdr_cells[1].text = "Capacidades"
    hdr_cells[2].text = "Contenidos"
    hdr_cells[3].text = "Criterios de evaluación"
    hdr_cells[4].text = "Evidencia de aprendizaje"
    hdr_cells[5].text = "Instrumento"
    for c in hdr_cells:
        c.paragraphs[0].runs[0].bold = True
        aplicar_fondo_celda(c, color_fondo_hex)
        
    tabla_area.cell(1, 0).text = propositos_data.get("competencias", "")
    tabla_area.cell(1, 1).text = propositos_data.get("capacidades", "")
    tabla_area.cell(1, 2).text = propositos_data.get("contenidos", "")
    tabla_area.cell(1, 3).text = propositos_data.get("criterios", "Hay que recordar que los criterios se desprenden de los estándares de aprendizaje.")
    tabla_area.cell(1, 4).text = propositos_data.get("evidencias", "")
    tabla_area.cell(1, 5).text = propositos_data.get("instrumentos", "")
    
    doc.add_paragraph()
    
    # Rúbrica
    tabla_rubrica = doc.add_table(rows=3, cols=5)
    tabla_rubrica.style = 'Table Grid'
    
    # Fila de Título
    titulo_rubrica = tabla_rubrica.cell(0, 0)
    titulo_rubrica.merge(tabla_rubrica.cell(0, 4))
    titulo_rubrica.text = "Rúbrica para evaluar el producto final de la sesión"
    titulo_rubrica.paragraphs[0].runs[0].bold = True
    titulo_rubrica.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    aplicar_fondo_celda(titulo_rubrica, color_fondo_hex)
    
    # Cabeceras Rúbrica
    hdr_cells = tabla_rubrica.rows[1].cells
    hdr_cells[0].text = "Criterios (alineados a la competencia y capacidades)"
    hdr_cells[1].text = "Inicio"
    hdr_cells[2].text = "En proceso"
    hdr_cells[3].text = "Logro esperado"
    hdr_cells[4].text = "Destacado"
    for c in hdr_cells:
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        aplicar_fondo_celda(c, color_fondo_hex)
        
    tabla_rubrica.cell(2, 0).text = propositos_data.get("criterios", "Criterios de evaluación de la sesión (redactar aquí cada criterio observable vinculado al producto final).")
    tabla_rubrica.cell(2, 1).text = "Desempeños iniciales: el producto de la sesión muestra logros muy parciales frente a los criterios (describir aquí)."
    tabla_rubrica.cell(2, 2).text = "Desempeños en desarrollo: el producto de la sesión cumple algunos aspectos de los criterios, pero aún presenta vacíos importantes."
    tabla_rubrica.cell(2, 3).text = "Logro esperado: el producto de la sesión cumple de manera consistente los criterios establecidos para la competencia y capacidades."
    tabla_rubrica.cell(2, 4).text = "Desempeño destacado: el producto de la sesión supera los criterios esperados, evidenciando mayor profundidad, creatividad o autonomía."

    doc.add_paragraph()
    
    # Valores
    tabla_valores = doc.add_table(rows=3, cols=4)
    tabla_valores.style = 'Table Grid'
    hdr_cells = tabla_valores.rows[0].cells
    hdr_cells[0].text = "Valor"
    hdr_cells[1].text = "Enfoques transversales"
    hdr_cells[2].text = "Comportamientos observables"
    hdr_cells[3].text = "Instrumento"
    for c in hdr_cells:
        c.paragraphs[0].runs[0].bold = True
        aplicar_fondo_celda(c, color_fondo_hex)
        
    celda_valor_f1 = tabla_valores.cell(1, 0)
    celda_valor_f1.merge(tabla_valores.cell(2, 0))
    celda_valor_f1.text = valores_enfoques_data.get("valores", "")
    
    tabla_valores.cell(1, 1).text = valores_enfoques_data.get("enfoques_transversales", "")
    tabla_valores.cell(1, 2).text = valores_enfoques_data.get("comportamientos_observables", "")
    tabla_valores.cell(1, 3).text = valores_enfoques_data.get("instrumento", "")

    doc.add_paragraph()
    
    # IV. SECUENCIA DE SESIONES
    doc.add_heading("IV. SECUENCIA DE SESIONES", 2)
    
    # Determinar filas necesarias de secuencias, 2 por fila
    sesiones_por_fila = 2
    num_sesiones = len(sesiones_data)
    num_filas = (num_sesiones + sesiones_por_fila - 1) // sesiones_por_fila
    
    tabla_secuencia = doc.add_table(rows=num_filas, cols=sesiones_por_fila)
    # Sin borde general
    
    for fila_idx in range(num_filas):
        for col_idx in range(sesiones_por_fila):
            sesion_num = fila_idx * sesiones_por_fila + col_idx + 1
            if sesion_num <= num_sesiones:
                datos = sesiones_data[sesion_num - 1]
                titulo = datos.get("titulo", "")
                titulo_limpio = re.sub(r'\\s*Duración:\\s*\d+\\s*(?:horas?|minutos?)\\.?\\s*$', '', (titulo or "").strip(), flags=re.IGNORECASE).strip()
                titulo_limpio = re.sub(r'\\s*\\(\d+\\s*(?:horas?|minutos?)\\)\\s*$', '', titulo_limpio, flags=re.IGNORECASE).strip()
                
                criterio = datos.get("criterio", "")
                actividades = datos.get("actividades", "")
                
                dur_valor = datos.get("duracion_valor", 4)
                dur_unidad = datos.get("duracion_unidad", "horas")
                duracion_texto = f"({int(float(dur_valor))} {dur_unidad})"
                
                def _quitar_duracion_texto(t):
                    if not t:
                        return t
                    s = re.sub(r'\\s*Duración:\\s*\d+\\s*(?:horas?|minutos?)\\.?\\s*', ' ', (t or ""), flags=re.IGNORECASE)
                    return re.sub(r'\\s+', ' ', s).strip()
                
                criterio_limpio = _quitar_duracion_texto(criterio)
                actividades_limpio = _quitar_duracion_texto(actividades)
                
                # Obtener la celda
                celda_sesion = tabla_secuencia.cell(fila_idx, col_idx)
                
                # Configurar celda padre para contener la tabla
                celda_sesion.text = ""
                # Si queremos evitar el párrafo extra al principio (es normal en python-docx)
                if celda_sesion.paragraphs:
                    p = celda_sesion.paragraphs[0]
                    p.text = ""
                
                # Crear tabla anidada dentro de la celda padre
                inner_table = celda_sesion.add_table(rows=4, cols=2)
                inner_table.style = 'Table Grid'
                
                # Header Sesion N
                hdr_inner = inner_table.cell(0, 0)
                hdr_inner.merge(inner_table.cell(0, 1))
                hdr_inner.text = f"Sesión N° {sesion_num}:   {duracion_texto}"
                hdr_inner.paragraphs[0].runs[0].bold = True
                aplicar_fondo_celda(hdr_inner, color_fondo_hex)
                
                # Titulo
                inner_table.cell(1, 0).text = "Título:"
                inner_table.cell(1, 0).paragraphs[0].runs[0].bold = True
                aplicar_fondo_celda(inner_table.cell(1, 0), color_fondo_hex)
                inner_table.cell(1, 1).text = titulo_limpio
                
                # Criterio
                inner_table.cell(2, 0).text = "Criterio de evaluación:"
                inner_table.cell(2, 0).paragraphs[0].runs[0].bold = True
                aplicar_fondo_celda(inner_table.cell(2, 0), color_fondo_hex)
                inner_table.cell(2, 1).text = criterio_limpio
                
                # Actividades
                inner_table.cell(3, 0).text = "Principales actividades de aprendizaje:"
                inner_table.cell(3, 0).paragraphs[0].runs[0].bold = True
                aplicar_fondo_celda(inner_table.cell(3, 0), color_fondo_hex)
                inner_table.cell(3, 1).text = actividades_limpio
                
                # Opcional: limpiar párrafos extra en la celda
                for paragraph in celda_sesion.paragraphs:
                    if not paragraph.text and not paragraph.runs:
                        p = paragraph._element
                        p.getparent().remove(p)
                        p._p = p._element = None
    
    # Pie de página
    doc.add_page_break()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_text = f"GENERADO POR SISTEMA IA EDUCATIVA\n"
    footer_text += f"Unidad: {titulo_unidad}\n"
    footer_text += "Ministerio de Educación - República del Perú"
    footer_run = footer.add_run(footer_text)
    footer_run.italic = True
    
    # Convertir a bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
