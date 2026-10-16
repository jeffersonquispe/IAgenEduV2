"""HU-23: exportación a DOCX con formato institucional, a partir de Document.content.

Antes, POST /export/docx volcaba cada sección como un encabezado genérico más
una tabla sin estilo. Aquí se construye un documento con la presentación de los
formatos oficiales: encabezado con la institución, título del documento,
secciones numeradas en el orden de la planificación, tablas con cabecera
sombreada, fichas (ITEM | CONTENIDO) con la primera columna resaltada, la
secuencia de sesiones como fichas por sesión y pie de página con numeración.

Reutiliza `aplicar_fondo_celda` y la paleta de src/app/docx_generator.py (la
exportación de la app Streamlit legacy). No reutiliza
`crear_documento_unidad_didactica`: recibe estructuras del parser legacy y
rellena la rúbrica con textos de ejemplo fijos, así que no sirve para un
documento ya generado; tampoco existe una versión para Sesión.
"""
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..app.docx_generator import aplicar_fondo_celda
from .control_documentario import construir_control
from .markdown_table import aplanar_tabla_degenerada, parse_markdown_table

COLOR_ENCABEZADO = "D9E1F2"        # mismo tono que docx_generator.py
COLOR_TITULO = RGBColor(0x1F, 0x38, 0x64)
FUENTE = "Arial"
ANCHO_UTIL_CM = 17.0               # A4 (21 cm) con márgenes de 2 cm

# Orden y títulos de las secciones por tipo de documento. Solo se numeran las
# que existen en el contenido.
ORDEN_UNIDAD: List[Tuple[str, str]] = [
    ("datos_informativos", "DATOS INFORMATIVOS"),
    ("situacion_significativa", "SITUACIÓN SIGNIFICATIVA"),
    ("propositos_de_aprendizaje", "PROPÓSITOS DE APRENDIZAJE"),
    ("valores_y_enfoques_transversales", "MATRIZ AXIOLÓGICA: VALORES Y ENFOQUES TRANSVERSALES"),
    ("secuencia_de_sesiones", "SECUENCIA DE SESIONES"),
    ("materiales_basicos", "MATERIALES BÁSICOS"),
    ("rubrica_unidad", "RÚBRICA ANALÍTICA DE LA UNIDAD"),
]
ORDEN_SESION: List[Tuple[str, str]] = [
    ("datos_informativos", "DATOS INFORMATIVOS"),
    ("reto", "RETO"),
    ("propositos_de_aprendizaje", "PROPÓSITOS DE APRENDIZAJE"),
    ("valores_y_enfoques_transversales", "MATRIZ AXIOLÓGICA: VALORES Y ENFOQUES TRANSVERSALES"),
    ("secuencia_didactica", "SECUENCIA DIDÁCTICA"),
    ("reflexion_sobre_la_actividad", "REFLEXIÓN SOBRE LA ACTIVIDAD"),
    ("rubrica_sesion", "RÚBRICA DE LA SESIÓN"),
]
TITULO_DOCUMENTO = {"unidad": "UNIDAD DIDÁCTICA", "sesion": "SESIÓN DE APRENDIZAJE"}

_ROMANOS = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]


def _romano(n: int) -> str:
    return _ROMANOS[n - 1] if 1 <= n <= len(_ROMANOS) else str(n)


# --- Utilidades de formato ---------------------------------------------------

def _fuente(run, size: float = 10, bold: bool = False, color: Optional[RGBColor] = None, italic: bool = False):
    run.font.name = FUENTE
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _escribir_celda(celda, texto: str, size: float = 9, bold: bool = False,
                    alinear: Optional[int] = None) -> None:
    """Texto en la celda; cada línea del texto es un párrafo."""
    lineas = [l.strip() for l in str(texto or "").split("\n") if l.strip()] or [""]
    for i, linea in enumerate(lineas):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if alinear is not None:
            p.alignment = alinear
        _fuente(p.add_run(linea), size=size, bold=bold)


def _repetir_encabezado(fila) -> None:
    """La fila de cabecera se repite al saltar de página."""
    trPr = fila._tr.get_or_add_trPr()
    marca = OxmlElement("w:tblHeader")
    marca.set(qn("w:val"), "true")
    trPr.append(marca)


def _fijar_anchos(tabla, proporciones: List[float]) -> None:
    total = sum(proporciones)
    tabla.autofit = False
    for fila in tabla.rows:
        for celda, p in zip(fila.cells, proporciones):
            celda.width = Cm(ANCHO_UTIL_CM * p / total)


def _campo(parrafo, instruccion: str) -> None:
    """Inserta un campo de Word (PAGE, NUMPAGES) en el párrafo."""
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, instruccion), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)
    _fuente(run, size=8)


# --- Estructura del documento ------------------------------------------------

def _configurar_documento(doc, titulo: str) -> None:
    for seccion in doc.sections:
        seccion.page_width, seccion.page_height = Cm(21.0), Cm(29.7)      # A4
        seccion.left_margin = seccion.right_margin = Cm(2.0)
        seccion.top_margin, seccion.bottom_margin = Cm(2.2), Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    doc.core_properties.title = titulo
    doc.core_properties.author = "IA-GEN-EDU"


def _celda_etiquetada(celda, etiqueta: str, valor: str = "", campo_paginas: bool = False) -> None:
    """'Etiqueta: valor' con la etiqueta en negrita (y, opcionalmente, 'N de M' páginas)."""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    _fuente(p.add_run(f"{etiqueta}: "), size=8, bold=True)
    if campo_paginas:
        _campo(p, "PAGE")
        _fuente(p.add_run(" de "), size=8)
        _campo(p, "NUMPAGES")
    elif valor:
        _fuente(p.add_run(valor), size=8)


def _encabezado_control_documentario(encabezado, control: Dict[str, str]) -> None:
    """HU-47: tabla de control documentario en el encabezado de página (se repite en cada hoja)."""
    tabla = encabezado.add_table(rows=2, cols=4, width=Cm(ANCHO_UTIL_CM))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    f1, f2 = tabla.rows[0].cells, tabla.rows[1].cells
    _celda_etiquetada(f1[0], "Documento", control["documento"])
    _celda_etiquetada(f1[1], "Institución", control["institucion"])
    _celda_etiquetada(f1[2], "Versión", control["version"])
    _celda_etiquetada(f1[3], "Páginas", campo_paginas=True)
    _celda_etiquetada(f2[0], "Aprobado por", control["aprobado_por"])
    _celda_etiquetada(f2[1], "Fecha", control["fecha_aprobacion"])
    _celda_etiquetada(f2[2], "Revisado por", control["revisado_por"])
    _celda_etiquetada(f2[3], "Fecha", control["fecha_revision"])
    _fijar_anchos(tabla, [4, 4, 2, 2])
    # Un encabezado debe terminar en un párrafo: el que ya existía queda detrás de la tabla.
    encabezado.paragraphs[0]._p.addprevious(tabla._tbl)


def _encabezado_y_pie(doc, tipo: str, perfil: Optional[Any]) -> None:
    seccion = doc.sections[0]
    control = construir_control(perfil, tipo)
    if control:
        # HU-47: con control documentario configurado, el encabezado es esa tabla.
        _encabezado_control_documentario(seccion.header, control)
    else:
        cab = seccion.header.paragraphs[0]
        cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
        nombre = (getattr(perfil, "nombre_institucion", None) or "").strip()
        ugel = (getattr(perfil, "dre_ugel", None) or "").strip()
        if nombre:
            _fuente(cab.add_run(nombre.upper()), size=11, bold=True, color=COLOR_TITULO)
            if ugel:
                cab.add_run("\n")
                _fuente(cab.add_run(ugel), size=8)
        else:
            _fuente(cab.add_run(TITULO_DOCUMENTO[tipo]), size=9, bold=True, color=COLOR_TITULO)

    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _fuente(pie.add_run("Página "), size=8)
    _campo(pie, "PAGE")
    _fuente(pie.add_run(" de "), size=8)
    _campo(pie, "NUMPAGES")
    pie.add_run("\n")
    _fuente(pie.add_run("Documento generado con IA-GEN-EDU. Revisar y ajustar antes de su uso."),
            size=7, italic=True, color=RGBColor(0x59, 0x59, 0x59))


def _titulo_documento(doc, tipo: str, titulo: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    _fuente(p.add_run(TITULO_DOCUMENTO[tipo]), size=16, bold=True, color=COLOR_TITULO)
    limpio = re.sub(r"^(Unidad Didáctica|Sesión)\s*-\s*", "", (titulo or "").strip())
    if limpio:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.space_after = Pt(10)
        _fuente(sub.add_run(limpio), size=12, bold=True)


def _titulo_seccion(doc, numero: int, texto: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _fuente(p.add_run(f"{_romano(numero)}. {texto}"), size=11, bold=True, color=COLOR_TITULO)


# --- Contenido de cada sección -----------------------------------------------

def _es_ficha(filas: List[List[str]]) -> bool:
    """Tabla ITEM | CONTENIDO de dos columnas: se muestra como ficha."""
    return (len(filas[0]) == 2 and filas[0][0].strip().upper() == "ITEM"
            and filas[0][1].strip().upper() == "CONTENIDO")


def _render_tabla(doc, filas: List[List[str]]) -> None:
    ficha = _es_ficha(filas)
    cuerpo = filas[1:] if ficha else filas
    ncols = len(filas[0])
    tabla = doc.add_table(rows=len(cuerpo), cols=ncols)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, fila in enumerate(cuerpo):
        for j in range(ncols):
            texto = fila[j] if j < len(fila) else ""
            celda = tabla.cell(i, j)
            texto = texto.replace("**", "")
            if ficha and j == 0:
                _escribir_celda(celda, texto, bold=True)
                aplicar_fondo_celda(celda, COLOR_ENCABEZADO)
            elif not ficha and i == 0:
                _escribir_celda(celda, texto, bold=True, alinear=WD_ALIGN_PARAGRAPH.CENTER)
                aplicar_fondo_celda(celda, COLOR_ENCABEZADO)
            else:
                _escribir_celda(celda, texto)
    if not ficha:
        _repetir_encabezado(tabla.rows[0])
    _fijar_anchos(tabla, [3, 7] if ficha else [1] * ncols)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


_PATRON_SESION = re.compile(r"^Sesi[oó]n\s+(\d+)\s*:\s*(.*)$", re.IGNORECASE)


def _bloques_de_sesion(texto: str) -> List[Dict[str, str]]:
    """Divide la secuencia de sesiones en [{numero, titulo, criterio, actividades}]."""
    sesiones: List[Dict[str, str]] = []
    actual: Optional[Dict[str, str]] = None
    campo = "titulo"
    for linea in texto.split("\n"):
        l = linea.strip()
        if not l:
            continue
        m = _PATRON_SESION.match(l)
        if m:
            actual = {"numero": m.group(1), "titulo": m.group(2).strip(), "criterio": "", "actividades": ""}
            sesiones.append(actual)
            campo = "titulo"
            continue
        if actual is None:
            return []
        if l.lower().startswith("criterio"):
            campo, l = "criterio", l.split(":", 1)[-1].strip()
        elif l.lower().startswith("actividades"):
            campo, l = "actividades", l.split(":", 1)[-1].strip()
        actual[campo] = (actual[campo] + " " + l).strip()
    return sesiones


def _render_secuencia_de_sesiones(doc, texto: str) -> bool:
    sesiones = _bloques_de_sesion(texto)
    if not sesiones:
        return False
    for s in sesiones:
        tabla = doc.add_table(rows=4, cols=2)
        tabla.style = "Table Grid"
        cab = tabla.cell(0, 0).merge(tabla.cell(0, 1))
        _escribir_celda(cab, f"Sesión N° {s['numero']}", bold=True, alinear=WD_ALIGN_PARAGRAPH.CENTER)
        aplicar_fondo_celda(cab, COLOR_ENCABEZADO)
        for i, (etiqueta, valor) in enumerate((
            ("Título", s["titulo"]), ("Criterio de evaluación", s["criterio"]),
            ("Principales actividades de aprendizaje", s["actividades"]),
        ), start=1):
            _escribir_celda(tabla.cell(i, 0), etiqueta, bold=True)
            aplicar_fondo_celda(tabla.cell(i, 0), COLOR_ENCABEZADO)
            _escribir_celda(tabla.cell(i, 1), valor)
        _fijar_anchos(tabla, [3, 7])
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return True


def _render_texto(doc, texto: str) -> None:
    parrafos = [p.strip() for p in re.split(r"\n\s*\n", str(texto)) if p.strip()] or [str(texto).strip()]
    for bloque in parrafos:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        lineas = [l.strip() for l in bloque.split("\n") if l.strip()]
        for i, linea in enumerate(lineas):
            if i:
                p.add_run().add_break()
            _fuente(p.add_run(linea.replace("**", "")), size=10)


def _render_seccion(doc, clave: str, valor: Any) -> None:
    texto = str(valor)
    if clave == "secuencia_de_sesiones" and _render_secuencia_de_sesiones(doc, texto):
        return
    filas = parse_markdown_table(texto)
    if filas:
        _render_tabla(doc, filas)
    else:
        # Una "tabla" con toda la información en una columna (p. ej. tras refinar con IA) va como texto.
        plano = aplanar_tabla_degenerada(texto)
        _render_texto(doc, plano if plano is not None else texto)


# --- API pública --------------------------------------------------------------

def _secciones_en_orden(data: Dict[str, Any], tipo: str) -> List[Tuple[str, str]]:
    orden = ORDEN_UNIDAD if tipo == "unidad" else ORDEN_SESION
    conocidas = {k for k, _ in orden}
    secciones = [(k, t) for k, t in orden if str(data.get(k) or "").strip()]
    # Cualquier otro campo con contenido se conserva al final: no se pierde nada.
    extras = [(k, k.replace("_", " ").upper()) for k, v in data.items()
              if k != "titulo" and k not in conocidas and str(v or "").strip()]
    return secciones + extras


def generar_docx(data: Dict[str, Any], tipo: str, perfil: Optional[Any] = None) -> bytes:
    """Devuelve el .docx de un documento ya generado.

    `data` es el contenido exportable (claves de Document.content), `tipo` es
    "unidad" o "sesion" y `perfil` el InstitutionProfile del docente (opcional).
    """
    if tipo not in TITULO_DOCUMENTO:
        raise ValueError(f"tipo de documento no soportado: {tipo!r}")
    doc = Document()
    titulo = str(data.get("titulo") or TITULO_DOCUMENTO[tipo].title())
    _configurar_documento(doc, titulo)
    _encabezado_y_pie(doc, tipo, perfil)
    _titulo_documento(doc, tipo, titulo)

    for numero, (clave, texto) in enumerate(_secciones_en_orden(data, tipo), start=1):
        _titulo_seccion(doc, numero, texto)
        _render_seccion(doc, clave, data[clave])

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
