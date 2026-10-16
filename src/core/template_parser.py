# src/core/template_parser.py
"""
Plantillas institucionales personalizables (HU-48/HU-49/HU-50).

Un docente puede subir el .docx "silueta" oficial de su institución para
Unidad Didáctica o Sesión de Aprendizaje (ver `Unidad didáctica(3).docx`,
`Sesión de Aprendizaje(1).docx` — versión en blanco del mismo formato que
`Unidad ejemplo.docx`/`Sesión de Aprendizaje ejm.docx`). Este módulo:

1. Parsea ese .docx en un esquema estructurado reutilizable (`parsear_plantilla`).
2. Construye, a partir de ese esquema, un prompt de generación dinámico que
   reemplaza los prompts estáticos `src/prompts/*.md` (`construir_prompt_desde_plantilla`).
3. Rellena el .docx original subido con el contenido ya generado, en vez de
   reconstruir el documento desde cero (`rellenar_plantilla_docx`).
"""
import io
import logging
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

import docx

from .markdown_table import parse_markdown_table

logger = logging.getLogger(__name__)

# Fragmentos de texto que, aunque no estén vacíos, delatan una celda de
# "instrucción para el docente" en la silueta en blanco (ver Tabla 3 de
# `Unidad didáctica(3).docx`: "Hay que recordar que los criterios se
# desprenden de los estándares de aprendizaje.") en vez de un dato fijo real
# como "Red Salesiana de Escuelas RSE". Heurística deliberadamente simple —
# HU-48 recomienda que el docente confirme/edite el esquema extraído antes
# de activarlo, en vez de confiar ciegamente en esto.
_MARCADORES_PLACEHOLDER = ("[", "…", "...", "……")
_MARCADORES_INSTRUCCION = (
    "hay que recordar", "pueden incluir", "ejemplo:", "se recomienda",
)


@dataclass
class TemplateCell:
    text: str
    row: int
    col: int
    es_dato: bool  # True = celda a completar por la IA; False = texto fijo del formato


@dataclass
class TemplateTable:
    index: int
    headers: List[str]
    filas: List[List[TemplateCell]] = field(default_factory=list)


@dataclass
class TemplateSchema:
    encabezados: List[str]
    tablas: List[TemplateTable]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "encabezados": self.encabezados,
            "tablas": [
                {
                    "index": t.index,
                    "headers": t.headers,
                    "filas": [
                        [
                            {"text": c.text, "row": c.row, "col": c.col, "es_dato": c.es_dato}
                            for c in fila
                        ]
                        for fila in t.filas
                    ],
                }
                for t in self.tablas
            ],
        }


def _es_celda_dato(texto: str, fila_idx: int) -> bool:
    """Heurística de HU-48: la fila 0 (encabezado de columnas) nunca es dato.
    Una celda vacía, con placeholder, o con texto instructivo para el
    docente se considera "a completar por la IA"; cualquier otro texto no
    vacío (encabezados fijos, valores institucionales ya escritos) se trata
    como contenido fijo que debe copiarse tal cual."""
    if fila_idx == 0:
        return False
    t = texto.strip()
    if not t:
        return True
    tl = t.lower()
    if any(m in t for m in _MARCADORES_PLACEHOLDER):
        return True
    if any(m in tl for m in _MARCADORES_INSTRUCCION):
        return True
    return False


def parsear_plantilla(contenido_bytes: bytes) -> TemplateSchema:
    """Parsea un .docx subido en un TemplateSchema. Lanza ValueError si el
    archivo no contiene ninguna tabla reconocible como plantilla de
    planificación (HU-48, criterio de aceptación de rechazo)."""
    doc = docx.Document(io.BytesIO(contenido_bytes))

    encabezados = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    tablas: List[TemplateTable] = []
    for idx, table in enumerate(doc.tables):
        if not table.rows:
            continue
        headers = [c.text.strip() for c in table.rows[0].cells]
        filas: List[List[TemplateCell]] = []
        for r, row in enumerate(table.rows):
            fila_celdas = [
                TemplateCell(
                    text=cell.text.strip(),
                    row=r,
                    col=c,
                    es_dato=_es_celda_dato(cell.text, r),
                )
                for c, cell in enumerate(row.cells)
            ]
            filas.append(fila_celdas)
        tablas.append(TemplateTable(index=idx, headers=headers, filas=filas))

    if not tablas:
        raise ValueError(
            "El documento no contiene ninguna tabla reconocible como plantilla "
            "de Unidad Didáctica o Sesión de Aprendizaje."
        )

    return TemplateSchema(encabezados=encabezados, tablas=tablas)


def schema_desde_dict(esquema: Dict[str, Any]) -> TemplateSchema:
    """Reconstruye un TemplateSchema desde el JSON guardado en DB (columna `esquema`)."""
    tablas = [
        TemplateTable(
            index=t["index"],
            headers=t.get("headers", []),
            filas=[[TemplateCell(**celda) for celda in fila] for fila in t.get("filas", [])],
        )
        for t in esquema.get("tablas", [])
    ]
    return TemplateSchema(encabezados=esquema.get("encabezados", []), tablas=tablas)


# ---------------------------------------------------------------------------
# HU-49: construir el prompt de generación a partir del esquema de plantilla
# ---------------------------------------------------------------------------

def construir_prompt_desde_plantilla(
    esquema: Dict[str, Any],
    instrucciones_pedagogicas: str,
    contexto: str,
) -> str:
    """Construye un prompt que le pide al modelo llenar EXACTAMENTE las
    tablas/columnas de la plantilla institucional activa, en vez de la
    estructura fija de unidad_didactica.md/sesion_aprendizaje.md.

    Cada celda de dato ("es_dato": true) se marca como `[GENERAR: <guía>]`;
    cada celda fija se marca como `[FIJO: <texto>]` y debe copiarse literal.
    """
    partes = [
        "Actúa como especialista en diseño curricular del MINEDU Perú. "
        "Debes generar el documento siguiendo EXACTAMENTE la estructura de "
        "tablas de la plantilla institucional a continuación — no agregues "
        "ni quites tablas, columnas ni filas fuera de las indicadas.",
        contexto,
        instrucciones_pedagogicas,
        "## ESTRUCTURA EXACTA A GENERAR (responde en Markdown, una tabla por bloque, en el mismo orden)",
    ]

    for encabezado in esquema.get("encabezados", []):
        partes.append(f"[Encabezado fijo del documento — inclúyelo tal cual, no lo regeneres]: {encabezado}")

    for tabla in esquema.get("tablas", []):
        headers = tabla.get("headers") or []
        filas = tabla.get("filas") or []
        partes.append(f"\nTABLA {tabla['index'] + 1}:")
        partes.append("| " + " | ".join(headers) + " |")
        partes.append("|" + "|".join(["---"] * max(len(headers), 1)) + "|")
        for fila in filas[1:]:  # fila 0 = encabezado de columnas, ya listado arriba
            celdas_texto = []
            for celda in fila:
                if celda.get("es_dato"):
                    guia = celda.get("text") or "genera el contenido apropiado para esta columna"
                    celdas_texto.append(f"[GENERAR: {guia}]")
                else:
                    celdas_texto.append(f"[FIJO: {celda.get('text', '')}]")
            partes.append("| " + " | ".join(celdas_texto) + " |")

    partes.append(
        "\nREGLAS:\n"
        "1. Reemplaza cada '[GENERAR: ...]' con contenido real, pedagógicamente correcto y "
        "coherente con el resto del documento.\n"
        "2. Cada '[FIJO: ...]' debe copiarse literalmente, sin modificar su texto.\n"
        "3. No agregues tablas, columnas ni filas adicionales a las indicadas arriba.\n"
        "4. Responde solo con las tablas en formato Markdown, sin texto antes ni después."
    )

    return "\n".join(p for p in partes if p)


# ---------------------------------------------------------------------------
# HU-50: rellenar el .docx original subido con el contenido ya generado
# ---------------------------------------------------------------------------

def _celdas_dato_en_orden(schema: TemplateSchema) -> List[Tuple[int, int, int]]:
    """Lista (tabla_idx, fila_idx, col_idx) de todas las celdas 'a completar', en orden."""
    objetivos = []
    for tabla in schema.tablas:
        for fila in tabla.filas:
            for celda in fila:
                if celda.es_dato:
                    objetivos.append((tabla.index, celda.row, celda.col))
    return objetivos


def _valores_generados_en_orden(contenido: Dict[str, Any]) -> List[str]:
    """Aplana el contenido generado (dict de secciones con texto/tablas Markdown)
    en una lista de valores de celda, en el mismo espíritu que las celdas de
    dato de la plantilla: si una sección es una tabla Markdown, se listan sus
    celdas (sin encabezado); si es texto plano, se agrega como un único valor.
    """
    valores: List[str] = []
    for clave, valor in contenido.items():
        if clave == "titulo" or not valor:
            continue
        tabla = parse_markdown_table(str(valor))
        if tabla and len(tabla) > 1:
            for fila in tabla[1:]:
                valores.extend(c for c in fila if c)
        else:
            valores.append(str(valor).strip())
    return valores


def rellenar_plantilla_docx(
    archivo_original: bytes,
    esquema: Dict[str, Any],
    contenido: Dict[str, Any],
) -> Tuple[bytes, Optional[str]]:
    """Abre el .docx original subido por el docente y escribe el contenido
    generado en sus celdas de dato, preservando diseño, encabezado y
    control documentario tal como el docente lo subió (HU-50).

    Es una asignación posicional heurística (celdas de dato de la plantilla,
    en orden, ← valores generados, en orden) — no una correspondencia
    semántica columna a columna. Si las cantidades no calzan, se rellena lo
    que se pueda y se devuelve una advertencia en vez de fallar o perder
    contenido silenciosamente.
    """
    schema = schema_desde_dict(esquema)
    objetivos = _celdas_dato_en_orden(schema)
    valores = _valores_generados_en_orden(contenido)

    doc = docx.Document(io.BytesIO(archivo_original))

    advertencia = None
    if len(valores) != len(objetivos):
        advertencia = (
            f"La plantilla tiene {len(objetivos)} celdas a completar pero se generaron "
            f"{len(valores)} valores; se rellenó hasta donde fue posible. Revisa el "
            f"documento exportado antes de compartirlo."
        )
        logger.warning("[TemplateFiller] %s", advertencia)

    n = min(len(objetivos), len(valores))
    for i in range(n):
        tabla_idx, fila_idx, col_idx = objetivos[i]
        try:
            doc.tables[tabla_idx].cell(fila_idx, col_idx).text = valores[i]
        except IndexError:
            continue

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue(), advertencia
