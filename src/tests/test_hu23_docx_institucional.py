"""HU-23: exportación DOCX con formato institucional (src/core/docx_export.py)."""
import io
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from docx import Document
from docx.shared import Cm

from src.core.docx_export import ORDEN_SESION, ORDEN_UNIDAD, generar_docx

PERFIL = SimpleNamespace(nombre_institucion="Colegio Salesiano de Cusco", dre_ugel="UGEL Cusco")

MATRIZ = (
    "| Valor | Enfoques transversales | Comportamientos observables | Instrumento |\n"
    "|---|---|---|---|\n"
    "| Fe | Bien común | Practica la oración. | Lista de cotejo |\n"
    "| Amor | Derechos | Llama por su nombre a los compañeros. | Escala de valoración |\n"
)
RUBRICA = (
    "| Criterio | C — Inicio | B — Proceso | A — Logrado | AD — Destacado |\n|---|---|---|---|---|\n"
    "| Identifica ideas | a | b | c | d |\n"
)
DATOS = (
    "| ITEM | CONTENIDO |\n|------|-----------|\n"
    "| **DRE / UGEL** | UGEL Cusco |\n| **Área curricular** | Comunicación |\n"
)

UNIDAD = {
    # Orden deliberadamente desordenado: el documento debe seguir el orden oficial.
    "rubrica_unidad": RUBRICA,
    "titulo": "Unidad Didáctica - Comunicación",
    "secuencia_de_sesiones": (
        "Sesión 1: La noticia\nCriterio: Identifica ideas\nActividades: Leer noticias\n\n"
        "Sesión 2: El reportaje\nCriterio: Compara textos\nActividades: Comparar"
    ),
    "valores_y_enfoques_transversales": MATRIZ,
    "situacion_significativa": "Los estudiantes enfrentan un reto real.\n\nSegundo párrafo.",
    "datos_informativos": DATOS,
    "propositos_de_aprendizaje": "Competencia: **Se comunica**.",
    "materiales_basicos": "Para el docente: proyector.\nPara el estudiante: cuaderno.",
}
SESION = {
    "titulo": "Sesión - La noticia",
    "datos_informativos": DATOS,
    "reto": "¿Cómo podemos informar a la comunidad?",
    "propositos_de_aprendizaje": "Competencia A",
    "secuencia_didactica": "INICIO 10 min. DESARROLLO 60 min. CIERRE 20 min.",
    "rubrica_sesion": "| Criterio | Logrado | En proceso |\n|---|---|---|\n| c1 | l | p |\n",
}


def _abrir(bytes_docx):
    return Document(io.BytesIO(bytes_docx))


def _parrafos(doc):
    return [p.text for p in doc.paragraphs]


def _relleno(celda):
    fills = celda._element.tcPr.xpath("w:shd/@w:fill") if celda._element.tcPr is not None else []
    return fills[0] if fills else None


def _texto_xml(parte):
    return parte._element.xml


# --- Estructura general ---

def test_es_un_docx_valido_con_formato_a4_y_metadatos():
    doc = _abrir(generar_docx(UNIDAD, "unidad", PERFIL))
    s = doc.sections[0]
    # Word guarda las medidas en twips: el redondeo difiere unos cientos de EMU (< 0,001 cm).
    assert abs(s.page_width - Cm(21.0)) < 1000 and abs(s.page_height - Cm(29.7)) < 1000
    assert abs(s.left_margin - Cm(2.0)) < 1000 and abs(s.right_margin - Cm(2.0)) < 1000
    assert doc.core_properties.title == "Unidad Didáctica - Comunicación"
    assert doc.styles["Normal"].font.name == "Arial"


def test_encabezado_con_la_institucion_y_la_ugel():
    doc = _abrir(generar_docx(UNIDAD, "unidad", PERFIL))
    cab = doc.sections[0].header.paragraphs[0].text
    assert "COLEGIO SALESIANO DE CUSCO" in cab and "UGEL Cusco" in cab


def test_sin_perfil_el_encabezado_muestra_el_tipo_de_documento():
    cab = _abrir(generar_docx(UNIDAD, "unidad", None)).sections[0].header.paragraphs[0].text
    assert cab == "UNIDAD DIDÁCTICA"
    cab = _abrir(generar_docx(SESION, "sesion", SimpleNamespace(nombre_institucion=None, dre_ugel=None))
                 ).sections[0].header.paragraphs[0].text
    assert cab == "SESIÓN DE APRENDIZAJE"


def test_pie_con_numeracion_de_paginas():
    xml = _texto_xml(_abrir(generar_docx(UNIDAD, "unidad", PERFIL)).sections[0].footer)
    assert "PAGE" in xml and "NUMPAGES" in xml
    assert "Documento generado con IA-GEN-EDU" in xml


def test_titulo_del_documento():
    p = _parrafos(_abrir(generar_docx(UNIDAD, "unidad", PERFIL)))
    assert p[0] == "UNIDAD DIDÁCTICA" and p[1] == "Comunicación"        # sin el prefijo "Unidad Didáctica - "


# --- Secciones: orden, numeración, contenido ---

def test_las_secciones_siguen_el_orden_oficial_y_se_numeran_en_romano():
    p = _parrafos(_abrir(generar_docx(UNIDAD, "unidad", PERFIL)))
    titulos = [t for t in p if t[:3] in ("I. ", "II.", "III", "IV.", "V. ", "VI.", "VII")]
    assert titulos == [
        "I. DATOS INFORMATIVOS", "II. SITUACIÓN SIGNIFICATIVA", "III. PROPÓSITOS DE APRENDIZAJE",
        "IV. MATRIZ AXIOLÓGICA: VALORES Y ENFOQUES TRANSVERSALES", "V. SECUENCIA DE SESIONES",
        "VI. MATERIALES BÁSICOS", "VII. RÚBRICA ANALÍTICA DE LA UNIDAD",
    ]


def test_la_numeracion_es_consecutiva_aunque_falten_secciones():
    parcial = {"titulo": "U", "situacion_significativa": "Algo", "rubrica_unidad": RUBRICA}
    p = _parrafos(_abrir(generar_docx(parcial, "unidad", None)))
    assert "I. SITUACIÓN SIGNIFICATIVA" in p and "II. RÚBRICA ANALÍTICA DE LA UNIDAD" in p


def test_las_secciones_vacias_no_aparecen():
    p = _parrafos(_abrir(generar_docx({**UNIDAD, "materiales_basicos": "   ", "rubrica_unidad": ""}, "unidad", None)))
    assert not any("MATERIALES" in t or "RÚBRICA" in t for t in p)


def test_los_campos_desconocidos_se_conservan_al_final():
    p = _parrafos(_abrir(generar_docx({**UNIDAD, "nota_extra": "Texto extra"}, "unidad", None)))
    assert p.index("VIII. NOTA EXTRA") > p.index("VII. RÚBRICA ANALÍTICA DE LA UNIDAD")
    assert "Texto extra" in p


def test_los_parrafos_y_las_negritas_markdown():
    p = _parrafos(_abrir(generar_docx(UNIDAD, "unidad", None)))
    assert "Los estudiantes enfrentan un reto real." in p and "Segundo párrafo." in p     # párrafos separados
    assert "Competencia: Se comunica." in p                                              # sin **


def test_la_sesion_tiene_su_propio_orden_y_titulo():
    doc = _abrir(generar_docx(SESION, "sesion", PERFIL))
    p = _parrafos(doc)
    assert p[0] == "SESIÓN DE APRENDIZAJE"
    assert p.index("II. RETO") > p.index("I. DATOS INFORMATIVOS")
    assert "¿Cómo podemos informar a la comunidad?" in p
    assert not any("SECUENCIA DE SESIONES" in t for t in p)


# --- Tablas ---

def test_datos_informativos_se_muestra_como_ficha_con_la_primera_columna_resaltada():
    doc = _abrir(generar_docx(UNIDAD, "unidad", PERFIL))
    ficha = doc.tables[0]
    assert len(ficha.rows) == 2 and len(ficha.columns) == 2          # sin la fila "ITEM | CONTENIDO"
    assert ficha.cell(0, 0).text == "DRE / UGEL" and ficha.cell(0, 1).text == "UGEL Cusco"
    assert _relleno(ficha.cell(0, 0)) == "D9E1F2" and _relleno(ficha.cell(0, 1)) is None


def test_la_matriz_axiologica_es_una_tabla_con_cabecera_sombreada_y_repetible():
    doc = _abrir(generar_docx(UNIDAD, "unidad", PERFIL))
    matriz = next(t for t in doc.tables if t.cell(0, 0).text == "Valor")
    assert [c.text for c in matriz.rows[0].cells] == [
        "Valor", "Enfoques transversales", "Comportamientos observables", "Instrumento"]
    assert all(_relleno(c) == "D9E1F2" for c in matriz.rows[0].cells)
    assert "w:tblHeader" in matriz.rows[0]._tr.xml                    # se repite al cambiar de página
    assert len(matriz.rows) == 3 and matriz.cell(2, 0).text == "Amor"


def test_la_rubrica_conserva_sus_cinco_columnas():
    doc = _abrir(generar_docx(UNIDAD, "unidad", PERFIL))
    rubrica = next(t for t in doc.tables if t.cell(0, 0).text == "Criterio")
    assert len(rubrica.columns) == 5 and rubrica.cell(1, 0).text == "Identifica ideas"


def test_la_secuencia_de_sesiones_se_muestra_como_una_ficha_por_sesion():
    doc = _abrir(generar_docx(UNIDAD, "unidad", PERFIL))
    fichas = [t for t in doc.tables if t.cell(0, 0).text.startswith("Sesión N°")]
    assert [f.cell(0, 0).text for f in fichas] == ["Sesión N° 1", "Sesión N° 2"]
    f1 = fichas[0]
    assert f1.cell(1, 1).text == "La noticia"
    assert f1.cell(2, 0).text == "Criterio de evaluación" and f1.cell(2, 1).text == "Identifica ideas"
    assert f1.cell(3, 1).text == "Leer noticias"


def test_una_secuencia_sin_el_patron_de_sesiones_cae_a_texto_sin_perder_contenido():
    doc = _abrir(generar_docx({**UNIDAD, "secuencia_de_sesiones": "Texto libre sin formato de sesiones."},
                              "unidad", None))
    assert "Texto libre sin formato de sesiones." in _parrafos(doc)


def test_tipo_invalido():
    with pytest.raises(ValueError):
        generar_docx(UNIDAD, "examen", None)


def test_el_orden_define_titulos_unicos():
    for orden in (ORDEN_UNIDAD, ORDEN_SESION):
        claves = [k for k, _ in orden]
        assert len(claves) == len(set(claves))


# --- Endpoint /export/docx (integración) ---

@pytest.mark.asyncio
async def test_endpoint_exporta_unidad_con_el_formato_institucional_y_el_perfil():
    from src.api.routes.export import export_docx

    db = MagicMock()
    # 1ª consulta: plantilla activa (ninguna); 2ª: perfil institucional.
    db.query.return_value.filter.return_value.first.side_effect = [None, PERFIL]
    user = SimpleNamespace(id="u1", stats=SimpleNamespace(exportaciones_docx=0))
    payload = {**UNIDAD, "done": True, "id": "x"}

    resp = await export_docx(data=payload, db=db, current_user=user)
    doc = Document(resp.path)

    assert "COLEGIO SALESIANO DE CUSCO" in doc.sections[0].header.paragraphs[0].text
    assert "I. DATOS INFORMATIVOS" in _parrafos(doc)
    assert any(t.cell(0, 0).text == "Valor" for t in doc.tables)
    assert "done" not in " ".join(_parrafos(doc)).lower().split()      # metadata del SSE no se exporta
    assert user.stats.exportaciones_docx == 1


@pytest.mark.asyncio
async def test_endpoint_distingue_sesion_de_unidad():
    from src.api.routes.export import export_docx

    db = MagicMock()
    db.query.return_value.filter.return_value.first.side_effect = [None, None]
    user = SimpleNamespace(id="u1", stats=None)
    resp = await export_docx(data=SESION, db=db, current_user=user)
    assert _parrafos(Document(resp.path))[0] == "SESIÓN DE APRENDIZAJE"
