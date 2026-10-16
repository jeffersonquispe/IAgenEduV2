"""HU-50 x HU-40/41: el llenado de la plantilla institucional con los campos
nuevos (reto, materiales_basicos). Usa un .docx real generado en memoria."""
import io

import docx

from src.core.template_parser import parsear_plantilla, rellenar_plantilla_docx


def _plantilla(celdas_vacias: int) -> bytes:
    d = docx.Document()
    t = d.add_table(rows=1 + celdas_vacias, cols=2)
    t.cell(0, 0).text = "ITEM"
    t.cell(0, 1).text = "CONTENIDO"
    for i in range(1, celdas_vacias + 1):
        t.cell(i, 0).text = f"Fijo {i}"
        t.cell(i, 1).text = ""          # celda a completar
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _celdas_col1(contenido: bytes):
    d = docx.Document(io.BytesIO(contenido))
    return [r.cells[1].text for r in d.tables[0].rows[1:]]


def test_plantilla_con_celda_para_reto_la_rellena_en_orden():
    archivo = _plantilla(3)
    esquema = parsear_plantilla(archivo).to_dict()
    contenido = {"titulo": "S", "datos_informativos": "Colegio X",
                 "reto": "¿Cómo lo resolvemos?", "propositos_de_aprendizaje": "Competencia A"}

    salida, advertencia = rellenar_plantilla_docx(archivo, esquema, contenido)

    assert advertencia is None
    assert _celdas_col1(salida) == ["Colegio X", "¿Cómo lo resolvemos?", "Competencia A"]


def test_plantilla_sin_celda_para_el_campo_nuevo_avisa_en_vez_de_perder_en_silencio():
    archivo = _plantilla(2)             # la plantilla no tiene lugar para el reto
    esquema = parsear_plantilla(archivo).to_dict()
    contenido = {"titulo": "S", "datos_informativos": "Colegio X",
                 "reto": "¿Cómo?", "propositos_de_aprendizaje": "Competencia A"}

    _, advertencia = rellenar_plantilla_docx(archivo, esquema, contenido)

    assert advertencia is not None and "3 valores" in advertencia


def test_materiales_basicos_de_unidad_se_escribe_en_la_plantilla():
    archivo = _plantilla(2)
    esquema = parsear_plantilla(archivo).to_dict()
    contenido = {"titulo": "U", "situacion_significativa": "Situación",
                 "materiales_basicos": "Docente: proyector\nEstudiante: cuaderno"}

    salida, advertencia = rellenar_plantilla_docx(archivo, esquema, contenido)

    assert advertencia is None
    assert "Estudiante: cuaderno" in _celdas_col1(salida)[1]
