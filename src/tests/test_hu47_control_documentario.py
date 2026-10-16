"""HU-47: encabezado de control documentario en PDF y DOCX, configurable desde el perfil."""
import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from src.api import models
from src.api.database import get_db
from src.api.routes import institution
from src.api.routes.auth import get_current_user
from src.core.control_documentario import CAMPOS, construir_control, normalizar_control
from src.core.docx_export import generar_docx
from src.core.pdf_generator import generar_pdf_sesion, generar_pdf_unidad

ROOT = Path(__file__).resolve().parents[2]
USER = "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa"

CONTROL = {
    "red": "Red Salesiana de Escuelas", "version": "02",
    "aprobado_por": "Dirección General", "fecha_aprobacion": "15/03/2026",
    "revisado_por": "Equipo pedagógico", "fecha_revision": "10/03/2026",
}


def _perfil(control=CONTROL, nombre="Colegio Salesiano de Cusco"):
    return SimpleNamespace(nombre_institucion=nombre, dre_ugel="UGEL Cusco", control_documentario=control)


CONTENIDO = {"titulo": "Unidad Didáctica - Comunicación", "situacion_significativa": "Una situación."}


# --- Datos del control ---

def test_normalizar_conserva_solo_campos_conocidos_y_no_vacios():
    assert normalizar_control({**CONTROL, "intruso": "x", "version": "  "}) == {
        k: v for k, v in CONTROL.items() if k != "version"}
    assert normalizar_control(None) == {} and normalizar_control("texto") == {} and normalizar_control({}) == {}


def test_normalizar_limpia_pipes_saltos_y_largos():
    r = normalizar_control({"aprobado_por": "Ana | Pérez\nGómez", "version": "1" * 50})
    assert r["aprobado_por"] == "Ana / Pérez Gómez"
    assert len(r["version"]) == CAMPOS["version"]


def test_sin_perfil_o_sin_control_no_hay_encabezado():
    assert construir_control(None, "unidad") is None
    assert construir_control(_perfil(control=None), "unidad") is None
    assert construir_control(_perfil(control={}), "unidad") is None
    assert construir_control(_perfil(control={"intruso": "x"}), "unidad") is None


def test_tipo_desconocido_no_genera_encabezado():
    assert construir_control(_perfil(), "examen") is None


def test_el_nombre_del_formato_depende_del_tipo():
    assert construir_control(_perfil(), "unidad")["documento"] == "Formato de Unidad Didáctica"
    assert construir_control(_perfil(), "sesion")["documento"] == "Formato de Sesión de Aprendizaje"


def test_la_red_tiene_prioridad_y_si_falta_se_usa_el_nombre_de_la_institucion():
    assert construir_control(_perfil(), "unidad")["institucion"] == "Red Salesiana de Escuelas"
    sin_red = {k: v for k, v in CONTROL.items() if k != "red"}
    assert construir_control(_perfil(control=sin_red), "unidad")["institucion"] == "Colegio Salesiano de Cusco"


def test_con_datos_parciales_el_resto_queda_en_blanco():
    c = construir_control(_perfil(control={"version": "03"}), "unidad")
    assert c["version"] == "03" and c["aprobado_por"] == "" and c["fecha_revision"] == ""


def test_si_la_columna_aun_no_existe_no_hay_encabezado_y_no_se_rompe():
    class SinColumna:
        nombre_institucion = "X"

        @property
        def control_documentario(self):
            raise RuntimeError("column institution_profiles.control_documentario does not exist")

    assert construir_control(SinColumna(), "unidad") is None


# --- DOCX ---

def _docx(data, tipo, perfil):
    return Document(io.BytesIO(generar_docx(data, tipo, perfil)))


def test_docx_con_control_pone_la_tabla_en_el_encabezado_de_pagina():
    cab = _docx(CONTENIDO, "unidad", _perfil()).sections[0].header
    assert len(cab.tables) == 1
    celdas = [[c.text for c in fila.cells] for fila in cab.tables[0].rows]
    assert celdas[0][:3] == ["Documento: Formato de Unidad Didáctica",
                             "Institución: Red Salesiana de Escuelas", "Versión: 02"]
    assert celdas[1] == ["Aprobado por: Dirección General", "Fecha: 15/03/2026",
                         "Revisado por: Equipo pedagógico", "Fecha: 10/03/2026"]


def test_docx_el_numero_de_paginas_es_un_campo_de_word():
    xml = _docx(CONTENIDO, "unidad", _perfil()).sections[0].header.tables[0].rows[0].cells[3]._element.xml
    assert "PAGE" in xml and "NUMPAGES" in xml
    assert _docx(CONTENIDO, "unidad", _perfil()).sections[0].header.tables[0].cell(0, 3).text.startswith("Páginas:")


def test_docx_el_encabezado_termina_en_un_parrafo_para_que_word_lo_abra():
    cab = _docx(CONTENIDO, "unidad", _perfil()).sections[0].header
    hijos = [el.tag.split("}")[1] for el in cab._element]
    assert hijos[-1] == "p" and "tbl" in hijos


def test_docx_sesion_usa_el_nombre_de_formato_de_sesion():
    cab = _docx({"titulo": "S", "reto": "¿Cómo?"}, "sesion", _perfil()).sections[0].header
    assert cab.tables[0].cell(0, 0).text == "Documento: Formato de Sesión de Aprendizaje"


def test_docx_sin_control_conserva_el_encabezado_anterior():
    for perfil in (None, _perfil(control=None), _perfil(control={})):
        cab = _docx(CONTENIDO, "unidad", perfil).sections[0].header
        assert not cab.tables                     # sin regresión: ningún encabezado nuevo
    assert "COLEGIO SALESIANO DE CUSCO" in _docx(CONTENIDO, "unidad", _perfil(control=None)
                                                  ).sections[0].header.paragraphs[0].text


def test_docx_el_cuerpo_no_cambia_con_el_control():
    con = [p.text for p in _docx(CONTENIDO, "unidad", _perfil()).paragraphs]
    sin = [p.text for p in _docx(CONTENIDO, "unidad", _perfil(control=None)).paragraphs]
    assert con == sin


# --- PDF ---

def _texto_pdf(ruta):
    return [p.extract_text() or "" for p in PdfReader(str(ruta)).pages]


def _control_pdf(tipo="unidad"):
    return construir_control(_perfil(), tipo)


def test_pdf_con_control_dibuja_el_encabezado_con_todos_los_datos(tmp_path):
    salida = tmp_path / "u.pdf"
    generar_pdf_unidad(CONTENIDO, str(salida), control=_control_pdf())
    texto = " ".join(_texto_pdf(salida))
    for esperado in ("Documento: Formato de Unidad Didáctica", "Institución: Red Salesiana de Escuelas",
                     "Versión: 02", "Aprobado por: Dirección General", "Fecha: 15/03/2026",
                     "Revisado por: Equipo pedagógico", "Fecha: 10/03/2026"):
        assert esperado in texto, esperado
    assert "PLANIFICACIÓN EDUCATIVA AI - MINEDU" not in texto          # lo reemplaza


def test_pdf_numera_paginas_como_n_de_m_y_repite_el_encabezado(tmp_path):
    largo = {"titulo": "U", "situacion_significativa": "\n\n".join(f"Párrafo {i}. " + "texto " * 80 for i in range(40))}
    salida = tmp_path / "largo.pdf"
    generar_pdf_unidad(largo, str(salida), control=_control_pdf())
    paginas = _texto_pdf(salida)
    assert len(paginas) >= 2
    total = len(paginas)
    for n, pagina in enumerate(paginas, 1):
        assert "Versión: 02" in pagina                                  # se repite en cada hoja
        assert f"Páginas: {n} de {total}" in pagina, (n, total)


def test_pdf_recorta_con_puntos_suspensivos_los_valores_demasiado_largos(tmp_path):
    control = {**_control_pdf(), "aprobado_por": "Una dirección con un nombre larguísimo " * 4}
    salida = tmp_path / "u.pdf"
    generar_pdf_unidad(CONTENIDO, str(salida), control=control)
    assert "…" in " ".join(_texto_pdf(salida))


def test_pdf_sesion_con_control(tmp_path):
    salida = tmp_path / "s.pdf"
    generar_pdf_sesion({"titulo": "S", "reto": "¿Cómo?"}, str(salida), control=_control_pdf("sesion"))
    assert "Formato de Sesión de Aprendizaje" in " ".join(_texto_pdf(salida))


def test_pdf_sin_control_no_cambia_el_encabezado_de_siempre(tmp_path):
    salida = tmp_path / "u.pdf"
    generar_pdf_unidad(CONTENIDO, str(salida))
    texto = " ".join(_texto_pdf(salida))
    assert "PLANIFICACIÓN EDUCATIVA AI - MINEDU" in texto and "Versión:" not in texto


# --- Rutas de exportación ---

@pytest.mark.asyncio
async def test_export_pdf_usa_el_control_del_perfil(tmp_path):
    from src.api.routes.export import export_pdf
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = _perfil()
    user = SimpleNamespace(id="u1", stats=None)
    resp = await export_pdf(data=CONTENIDO, db=db, current_user=user)
    assert "Versión: 02" in " ".join(_texto_pdf(resp.path))


@pytest.mark.asyncio
async def test_export_pdf_sin_perfil_conserva_el_encabezado_de_siempre():
    from src.api.routes.export import export_pdf
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = None
    resp = await export_pdf(data=CONTENIDO, db=db, current_user=SimpleNamespace(id="u1", stats=None))
    assert "PLANIFICACIÓN EDUCATIVA AI - MINEDU" in " ".join(_texto_pdf(resp.path))


@pytest.mark.asyncio
async def test_export_docx_usa_el_control_del_perfil():
    from src.api.routes.export import export_docx
    db = MagicMock()
    db.query.return_value.filter.return_value.first.side_effect = [None, _perfil()]    # plantilla, perfil
    resp = await export_docx(data=CONTENIDO, db=db, current_user=SimpleNamespace(id="u1", stats=None))
    assert Document(resp.path).sections[0].header.tables[0].cell(0, 2).text == "Versión: 02"


# --- Rutas del perfil: guardar y leer el control ---

def _cliente(tmp_path, con_columna=True):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    models.User.__table__.create(engine, checkfirst=True)
    if con_columna:
        models.InstitutionProfile.__table__.create(engine, checkfirst=True)
    else:
        # Base "anterior a la migración": la misma tabla, sin control_documentario.
        with engine.begin() as c:
            c.execute(text(
                "CREATE TABLE institution_profiles (id INTEGER PRIMARY KEY, user_id VARCHAR, "
                "nombre_institucion VARCHAR, nivel_modalidad VARCHAR, dre_ugel VARCHAR, "
                "coordinacion_pedagogica VARCHAR, vision VARCHAR, mision VARCHAR, contexto_relevante VARCHAR, "
                "valores_institucionales JSON, enfoques_transversales JSON, matriz_axiologica JSON, "
                "created_at DATETIME, updated_at DATETIME)"))
    sesion = sessionmaker(bind=engine)()
    sesion.add(models.User(id=USER, email="a@x.com"))
    sesion.commit()
    app = FastAPI()
    app.include_router(institution.router)
    app.dependency_overrides[get_db] = lambda: sesion
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=USER)
    return TestClient(app), sesion


def test_put_guarda_el_control_limpio_y_get_lo_devuelve(tmp_path):
    cliente, _ = _cliente(tmp_path)
    r = cliente.put("/institucion", json={"nombre_institucion": "X", "control_documentario": {
        "version": " 02 ", "aprobado_por": "Ana | Pérez", "revisado_por": ""}})
    assert r.status_code == 200
    assert r.json()["control_documentario"]["version"] == "02"
    assert r.json()["control_documentario"]["aprobado_por"] == "Ana / Pérez"
    guardado = cliente.get("/institucion").json()["control_documentario"]
    assert guardado["version"] == "02" and guardado["revisado_por"] is None      # lo vacío no se guarda


def test_vaciar_todos_los_campos_borra_el_control(tmp_path):
    cliente, _ = _cliente(tmp_path)
    cliente.put("/institucion", json={"control_documentario": {"version": "02"}})
    cliente.put("/institucion", json={"control_documentario": {"version": "", "aprobado_por": "  "}})
    assert cliente.get("/institucion").json()["control_documentario"] is None


def test_guardar_el_perfil_sin_mencionar_el_control_no_lo_borra(tmp_path):
    cliente, _ = _cliente(tmp_path)
    cliente.put("/institucion", json={"control_documentario": {"version": "02"}})
    cliente.put("/institucion", json={"vision": "Nueva visión"})
    assert cliente.get("/institucion").json()["control_documentario"]["version"] == "02"


@pytest.mark.parametrize("campo,valor", [("version", "x" * 21), ("aprobado_por", "x" * 81), ("fecha_revision", "x" * 41)])
def test_el_servidor_rechaza_valores_demasiado_largos(tmp_path, campo, valor):
    cliente, _ = _cliente(tmp_path)
    assert cliente.put("/institucion", json={"control_documentario": {campo: valor}}).status_code == 422


def test_antes_de_la_migracion_la_pantalla_de_institucion_sigue_funcionando(tmp_path):
    """Sin la columna en la base: GET responde (control = None) y se puede guardar el perfil."""
    cliente, sesion = _cliente(tmp_path, con_columna=False)
    r = cliente.put("/institucion", json={"nombre_institucion": "Colegio X", "vision": "V"})
    assert r.status_code == 200 and r.json()["control_documentario"] is None
    g = cliente.get("/institucion")
    assert g.status_code == 200 and g.json()["nombre_institucion"] == "Colegio X"
    assert g.json()["control_documentario"] is None


def test_antes_de_la_migracion_las_exportaciones_siguen_funcionando(tmp_path):
    _, sesion = _cliente(tmp_path, con_columna=False)
    sesion.add(models.InstitutionProfile(user_id=USER, nombre_institucion="Colegio X"))
    sesion.commit()
    perfil = sesion.query(models.InstitutionProfile).filter_by(user_id=USER).first()
    assert construir_control(perfil, "unidad") is None                   # no lanza
    assert _docx(CONTENIDO, "unidad", perfil).sections[0].header.tables == []


# --- Migración y frontend ---

def test_la_migracion_agrega_la_columna_de_forma_idempotente():
    sql = (ROOT / "supabase/migrations/20261009_institution_control_documentario.sql").read_text(encoding="utf-8")
    assert "ADD COLUMN IF NOT EXISTS control_documentario json" in sql and "BEGIN;" in sql and "COMMIT;" in sql


def test_la_pantalla_de_institucion_tiene_los_campos_y_envia_el_control_solo_si_se_usa():
    src = (ROOT / "frontend/src/pages/Institucion.jsx").read_text(encoding="utf-8")
    for nombre in CAMPOS:
        assert f'name="{nombre}"' in src, nombre
    assert "teniaControl.current" in src and "payload.control_documentario = control" in src
