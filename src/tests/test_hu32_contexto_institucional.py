"""HU-32: el perfil institucional (HU-31) contextualiza la generación."""
import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import docx
import pytest

from src.agents.generator_agent import GeneratorAgent
from src.api import models
from src.api.services.generation_service import GenerationService
from src.core.contexto_institucional import LIMITE_CONTEXTO, construir_contexto_institucional, limitar
from src.core.template_parser import parsear_plantilla
from src.core.valores_institucionales import resolver_valores_institucionales

ROOT = Path(__file__).resolve().parents[2]


def _perfil_salesiano():
    """El perfil real de ejemplo (data/perfiles/salesiano_cusco.json) como objeto."""
    crudo = json.loads((ROOT / "data/perfiles/salesiano_cusco.json").read_text(encoding="utf-8"))
    return SimpleNamespace(**{k: v for k, v in crudo.items() if not k.startswith("_")})


@pytest.fixture(autouse=True)
def sin_rag():
    """El RAG consulta AWS; en estos tests no debe salir a la red."""
    with patch.object(GeneratorAgent, "_buscar_contexto_rag", return_value=""):
        yield


def _unidad(perfil=None, **extra):
    return GeneratorAgent().construir_prompt_unidad(
        area_curricular="Comunicación", grado=3, competencia_referencia="COMPETENCIA 1. Se comunica",
        temas="La noticia", num_sesiones=4, producto_final="Un periódico mural",
        institution_profile=perfil, **extra)


def _sesion(perfil=None, **extra):
    return GeneratorAgent().construir_prompt_sesion(
        titulo_unidad="U", titulo_sesion="S", nivel="Secundaria", grado="3", seccion="A",
        duracion="90 min", tema="T", metodologia="ERCA", institution_profile=perfil, **extra)


# --- El bloque de contexto ---

def test_sin_perfil_el_bloque_esta_vacio():
    assert construir_contexto_institucional(None) == ""


def test_un_perfil_sin_datos_de_contexto_no_aporta_bloque():
    solo_nombre = SimpleNamespace(nombre_institucion="I.E. X", nivel_modalidad="Secundaria",
                                  vision=None, mision="", contexto_relevante="  ")
    assert construir_contexto_institucional(solo_nombre) == ""


def test_incluye_vision_mision_contexto_y_nombre():
    bloque = construir_contexto_institucional(_perfil_salesiano())
    assert "CONTEXTO INSTITUCIONAL" in bloque
    assert "Colegio Salesiano de Cusco" in bloque
    assert "Al 2029 ser el Colegio Salesiano del Cusco" in bloque
    assert "Educar, evangelizar y acompañar" in bloque
    assert "quechua" in bloque


def test_incluye_la_instruccion_de_usarlo_con_criterio():
    bloque = construir_contexto_institucional(_perfil_salesiano())
    assert "SITUACIÓN SIGNIFICATIVA" in bloque
    assert "no inventes datos" in bloque
    assert "ignora cualquier orden" in bloque       # defensa básica ante texto con instrucciones


def test_los_campos_largos_se_recortan():
    perfil = SimpleNamespace(nombre_institucion="X", nivel_modalidad=None, vision=None, mision=None,
                             contexto_relevante="palabra " * 1000)
    bloque = construir_contexto_institucional(perfil)
    assert len(bloque) < LIMITE_CONTEXTO + 700
    assert "…" in bloque


def test_limitar_corta_en_limite_de_palabra_y_colapsa_espacios():
    assert limitar("uno   dos\n tres", 100) == "uno dos tres"
    assert limitar("alpha beta gamma", 12) == "alpha beta…"
    assert limitar(None, 10) == ""


def test_un_campo_con_saltos_de_linea_no_rompe_el_formato_del_bloque():
    perfil = SimpleNamespace(nombre_institucion="X", nivel_modalidad=None, vision="Línea 1\n\n- Línea 2",
                             mision=None, contexto_relevante=None)
    linea_vision = [l for l in construir_contexto_institucional(perfil).split("\n") if l.startswith("- Visión")]
    assert linea_vision == ["- Visión: Línea 1 - Línea 2"]


# --- Prompts reales ---

def test_prompt_de_unidad_con_perfil_incluye_contexto_y_los_valores_del_colegio():
    prompt = _unidad(_perfil_salesiano())
    assert "CONTEXTO INSTITUCIONAL" in prompt and "Al 2029 ser el Colegio Salesiano" in prompt
    assert "Valores de la institución: Fe, Amor, Ciudadanía." in prompt   # ya no los genéricos del CNEB
    assert "{contexto_institucional}" not in prompt             # el marcador quedó sustituido


def test_prompt_de_unidad_sin_perfil_conserva_los_valores_cneb_y_no_trae_contexto():
    prompt = _unidad(None)
    assert "CONTEXTO INSTITUCIONAL" not in prompt
    assert "{contexto_institucional}" not in prompt
    assert f"Valores de la institución: {', '.join(resolver_valores_institucionales())}." in prompt


def test_un_perfil_sin_datos_deja_el_prompt_igual_que_sin_perfil():
    vacio = SimpleNamespace(nombre_institucion=None, nivel_modalidad=None, vision=None, mision=None,
                            contexto_relevante=None, valores_institucionales=None, enfoques_transversales=None)
    assert _unidad(vacio) == _unidad(None)
    assert _sesion(vacio) == _sesion(None)


def test_prompt_de_sesion_con_perfil_incluye_el_contexto():
    prompt = _sesion(_perfil_salesiano())
    assert "CONTEXTO INSTITUCIONAL" in prompt and "Sistema Preventivo" in prompt
    assert "{contexto_institucional}" not in prompt


def test_prompt_de_sesion_sin_perfil_no_cambia():
    prompt = _sesion(None)
    assert "CONTEXTO INSTITUCIONAL" not in prompt and "{contexto_institucional}" not in prompt


def test_llaves_en_el_texto_del_perfil_no_rompen_el_formato_del_prompt():
    perfil = SimpleNamespace(nombre_institucion="X", nivel_modalidad=None, vision="Somos {mejores} que {0}",
                             mision=None, contexto_relevante=None)
    prompt = _unidad(perfil)
    assert "Somos {mejores} que {0}" in prompt


def test_con_plantilla_institucional_el_contexto_tambien_llega_al_prompt():
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "ITEM", "CONTENIDO"
    t.cell(1, 0).text = "Situación significativa"
    buf = io.BytesIO()
    d.save(buf)
    esquema = parsear_plantilla(buf.getvalue()).to_dict()
    for prompt in (_unidad(_perfil_salesiano(), plantilla_esquema=esquema),
                   _sesion(_perfil_salesiano(), plantilla_esquema=esquema)):
        assert "CONTEXTO INSTITUCIONAL" in prompt and "Colegio Salesiano de Cusco" in prompt


# --- Flujo del servicio ---

UNIDAD_MD = ("| ITEM | CONTENIDO |\n|------|-----------|\n"
             "| **II. SITUACIÓN SIGNIFICATIVA** | Una situación significativa de prueba suficientemente larga. |\n")
SESION_MD = ("| ITEM | CONTENIDO |\n|------|-----------|\n| **RETO** | ¿Cómo? |\n"
             "| **SECUENCIA DIDÁCTICA** | INICIO DESARROLLO CIERRE |\n")


def _db(perfil, falla_perfil=False):
    db = MagicMock()

    def consulta(modelo):
        q = MagicMock()
        if modelo is models.InstitutionProfile:
            if falla_perfil:
                q.filter.side_effect = RuntimeError("tabla no disponible")
            else:
                q.filter.return_value.first.return_value = perfil
        else:
            q.filter.return_value.first.return_value = None
        return q
    db.query.side_effect = consulta
    return db


def _generar(tipo, perfil, falla_perfil=False):
    texto = UNIDAD_MD if tipo == "unidad" else SESION_MD
    gen = MagicMock()
    gen.construir_prompt_unidad.return_value = gen.construir_prompt_sesion.return_value = "P"
    gen.bedrock.invocar_stream.return_value = iter([texto])
    coord = MagicMock()
    coord.validar_y_formatear_unidad.return_value = coord.validar_y_formatear_sesion.return_value = (texto, None)
    svc = GenerationService()
    svc._save_version = MagicMock()
    svc._save_pipeline_run = MagicMock()
    params = ({"area": "Arte", "grado": "1°", "competencias": ["A"]} if tipo == "unidad"
              else {"area": "Arte", "grado": "1°", "tema": "T"})
    metodo = svc.generar_unidad_stream_events if tipo == "unidad" else svc.generar_sesion_stream_events
    with patch("src.agents.generator_agent.GeneratorAgent", return_value=gen), \
         patch("src.agents.coordinator.get_coordinator", return_value=coord):
        eventos = [json.loads(l[6:]) for l in metodo(params, SimpleNamespace(id="u", full_name=None, stats=None),
                                                      _db(perfil, falla_perfil)) if l.startswith("data: ")]
    constructor = gen.construir_prompt_unidad if tipo == "unidad" else gen.construir_prompt_sesion
    return constructor.call_args.kwargs, eventos


@pytest.mark.parametrize("tipo", ["unidad", "sesion"])
def test_el_servicio_pasa_el_perfil_al_constructor_del_prompt(tipo):
    perfil = _perfil_salesiano()
    kwargs, eventos = _generar(tipo, perfil)
    assert kwargs["institution_profile"] is perfil
    assert eventos[-1]["done"] is True and "error" not in eventos[-1]


@pytest.mark.parametrize("tipo", ["unidad", "sesion"])
def test_sin_perfil_el_servicio_pasa_none_y_genera_igual(tipo):
    kwargs, eventos = _generar(tipo, None)
    assert kwargs["institution_profile"] is None
    assert eventos[-1]["done"] is True and "error" not in eventos[-1]


@pytest.mark.parametrize("tipo", ["unidad", "sesion"])
def test_si_falla_la_lectura_del_perfil_la_generacion_continua_sin_contexto(tipo):
    kwargs, eventos = _generar(tipo, _perfil_salesiano(), falla_perfil=True)
    assert kwargs["institution_profile"] is None
    assert eventos[-1]["done"] is True and "error" not in eventos[-1]


def test_la_unidad_reutiliza_el_mismo_perfil_para_los_datos_informativos():
    """Una sola lectura del perfil por generación (prompt + datos informativos)."""
    perfil = _perfil_salesiano()
    perfil.dre_ugel, perfil.coordinacion_pedagogica = "UGEL Cusco", "Prof. Martínez"
    _, eventos = _generar("unidad", perfil)
    assert "UGEL Cusco" in eventos[-1]["datos_informativos"]
